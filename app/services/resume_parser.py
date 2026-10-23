"""Resume document extraction, normalization, and structured profile parsing."""

from __future__ import annotations

import re
import unicodedata
import zipfile
from dataclasses import dataclass
from pathlib import Path

import pdfplumber
from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from pdfminer.pdfdocument import PDFPasswordIncorrect
from pdfminer.pdfparser import PDFSyntaxError

from app.services.section_detector import detect_sections
from app.services.skill_extractor import extract_skills

EMAIL_PATTERN = re.compile(r"\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b", re.IGNORECASE)
PHONE_PATTERN = re.compile(r"(?<!\w)\+?\d[\d\s().-]{7,}\d(?!\w)")
LINK_PATTERN = re.compile(
    r"(?:https?://)?(?:www\.)?(?:linkedin\.com|github\.com)/[^\s|,;]+",
    re.IGNORECASE,
)
BULLET_PATTERN = re.compile(r"^[\s]*[•●▪◦‣∙]\s*")
WORD_HYPHENATION_PATTERN = re.compile(r"(?<=\w)-\s*\n\s*(?=\w)")
SPACE_PATTERN = re.compile(r"[^\S\n]+")
BLANK_LINES_PATTERN = re.compile(r"\n{3,}")


class ResumeParseError(ValueError):
    """A resume is unreadable or contains no extractable text."""


@dataclass(frozen=True)
class ParsedResume:
    """Normalized resume text and its extracted structured profile."""

    text: str
    profile: dict[str, object]


def normalize_resume_text(text: str) -> str:
    """Normalize unicode, bullets, whitespace, and line-broken words."""
    normalized = unicodedata.normalize("NFKC", text).replace("\u00ad", "")
    normalized = normalized.replace("\r\n", "\n").replace("\r", "\n")
    normalized = WORD_HYPHENATION_PATTERN.sub("", normalized)
    lines = []
    for line in normalized.splitlines():
        cleaned = SPACE_PATTERN.sub(" ", line).strip()
        cleaned = BULLET_PATTERN.sub("- ", cleaned)
        lines.append(cleaned)
    normalized = "\n".join(lines)
    normalized = BLANK_LINES_PATTERN.sub("\n\n", normalized)
    if not normalized.strip():
        raise ResumeParseError("No readable text was found in this resume.")
    return normalized.strip()


def extract_text_from_file(path: str | Path) -> str:
    """Extract selectable text from a PDF or DOCX file."""
    source = Path(path)
    extension = source.suffix.casefold()
    if extension == ".pdf":
        try:
            with pdfplumber.open(source) as pdf:
                pages = [page.extract_text() or "" for page in pdf.pages]
        except (OSError, ValueError, PDFSyntaxError, PDFPasswordIncorrect) as error:
            raise ResumeParseError(
                "This PDF could not be read. Please upload a valid, unlocked PDF."
            ) from error
        text = "\n".join(pages)
        if not text.strip():
            raise ResumeParseError(
                "No selectable text was found. Scanned or image-only PDFs are not supported."
            )
        return normalize_resume_text(text)

    if extension == ".docx":
        try:
            document = Document(source)
        except (OSError, ValueError, zipfile.BadZipFile, PackageNotFoundError) as error:
            raise ResumeParseError(
                "This DOCX file could not be read. Please upload a valid Word document."
            ) from error
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            paragraphs.extend(
                " | ".join(cell.text for cell in row.cells) for row in table.rows
            )
        return normalize_resume_text("\n".join(paragraphs))

    raise ResumeParseError("Unsupported resume format. Upload a PDF or DOCX file.")


def parse_resume_text(text: str) -> ParsedResume:
    """Normalize resume text and extract contact details and common sections."""
    normalized = normalize_resume_text(text)
    sections = detect_sections(normalized)
    email_match = EMAIL_PATTERN.search(normalized)
    phone_match = next(
        (
            match
            for match in PHONE_PATTERN.finditer(normalized)
            if 10 <= len(re.sub(r"\D", "", match.group())) <= 15
        ),
        None,
    )
    links = _extract_links(normalized)
    profile: dict[str, object] = {
        "name": _extract_name(normalized, email_match.group() if email_match else None),
        "email": email_match.group() if email_match else None,
        "phone": phone_match.group().strip() if phone_match else None,
        "links": links,
        "education": _section_entries(sections, "education"),
        "skills": extract_skills(normalized),
        "projects": _section_entries(sections, "projects"),
        "experience": _section_entries(sections, "experience"),
        "certifications": _section_entries(sections, "certifications"),
        "achievements": _section_entries(sections, "achievements"),
        "summary": sections.get("summary", ""),
        "sections": sorted(sections),
    }
    return ParsedResume(text=normalized, profile=profile)


def parse_resume_file(path: str | Path) -> ParsedResume:
    """Extract and parse a resume from a PDF or DOCX path."""
    return parse_resume_text(extract_text_from_file(path))


def _extract_links(text: str) -> dict[str, list[str]]:
    links = {"linkedin": [], "github": []}
    for match in LINK_PATTERN.finditer(text):
        url = match.group().rstrip(").]")
        service = "linkedin" if "linkedin.com" in url.casefold() else "github"
        links[service].append(url)
    return links


def _extract_name(text: str, email: str | None) -> str | None:
    """Use the first plausible header line, without mistaking contact details."""
    ignored = {
        "resume",
        "curriculum vitae",
        "cv",
        "profile",
        "summary",
        "education",
        "skills",
        "projects",
        "experience",
    }
    for line in text.splitlines()[:8]:
        candidate = line.strip().lstrip("- ").strip()
        if (
            not candidate
            or candidate.casefold() in ignored
            or (email and email in candidate)
            or PHONE_PATTERN.search(candidate)
            or LINK_PATTERN.search(candidate)
            or len(candidate) > 80
        ):
            continue
        return candidate
    return None


def _section_entries(sections: dict[str, str], name: str) -> list[str]:
    return [
        line.lstrip("- ").strip()
        for line in sections.get(name, "").splitlines()
        if line.strip()
    ]
