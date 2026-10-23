"""Authenticated resume upload, ownership, and persistence tests."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

from docx import Document

from app.extensions import db
from app.models import Resume
from app.services.resume_parser import parse_resume_file

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _register(client, email: str) -> None:
    response = client.post(
        "/register",
        data={
            "name": "Aarav Mehta",
            "email": email,
            "password": "strong-password",
            "confirm_password": "strong-password",
        },
        follow_redirects=True,
    )
    assert response.status_code == 200


def _make_docx(text: str) -> BytesIO:
    document = Document()
    for line in text.splitlines():
        document.add_paragraph(line)
    output = BytesIO()
    document.save(output)
    output.seek(0)
    return output


def _make_pdf(text: str) -> bytes:
    def escape_pdf_text(line: str) -> str:
        return line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")

    drawing = ["BT /F1 10 Tf 50 760 Td"]
    for line in text.splitlines():
        if line:
            drawing.append(f"({escape_pdf_text(line)}) Tj")
        drawing.append("0 -14 Td")
    drawing.append("ET")
    content = "\n".join(drawing).encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>"
        ),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(content)).encode("ascii") + b" >>\nstream\n"
        + content
        + b"\nendstream",
    ]
    pdf = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, body in enumerate(objects, start=1):
        offsets.append(len(pdf))
        pdf.extend(f"{number} 0 obj\n".encode("ascii"))
        pdf.extend(body)
        pdf.extend(b"\nendobj\n")
    xref_offset = len(pdf)
    pdf.extend(f"xref\n0 {len(offsets)}\n".encode("ascii"))
    pdf.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        pdf.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    pdf.extend(
        f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\n"
        f"startxref\n{xref_offset}\n%%EOF\n".encode("ascii")
    )
    return bytes(pdf)


def test_docx_upload_saves_parsed_profile_and_removes_original(client, app, sample_resume_text):
    _register(client, "aarav@example.com")

    response = client.post(
        "/resumes/",
        data={
            "resume_file": (
                _make_docx(sample_resume_text),
                "../aarav_resume.docx",
                DOCX_MIME,
            )
        },
        content_type="multipart/form-data",
        follow_redirects=True,
    )

    assert response.status_code == 200
    assert b"Resume uploaded and analyzed" in response.data
    assert b"Aarav Mehta" in response.data
    assert b"Natural Language Processing" not in response.data
    assert b"PostgreSQL" in response.data
    assert b"Campus Labs" in response.data
    assert list(Path(app.config["UPLOAD_FOLDER"]).glob("*")) == []
    with app.app_context():
        resume = db.session.scalar(db.select(Resume))
        assert resume is not None
        assert resume.original_filename == "aarav_resume.docx"
        assert resume.parsed_text is not None
        assert resume.profile_json["skills"]["Programming Languages"] == [
            "JavaScript",
            "Python",
        ]
        assert resume.is_active


def test_pdf_upload_parses_text_and_detects_skills(client, sample_resume_text):
    _register(client, "pdf@example.com")

    response = client.post(
        "/resumes/",
        data={
            "resume_file": (
                BytesIO(_make_pdf(sample_resume_text)),
                "aarav.pdf",
                "application/pdf",
            )
        },
        content_type="multipart/form-data",
        follow_redirects=True,
    )

    assert response.status_code == 200
    assert b"Resume uploaded and analyzed" in response.data
    assert b"aarav.mehta@example.com" in response.data
    assert b"Python" in response.data


def test_parser_reads_docx_and_pdf_files(tmp_path, sample_resume_text):
    docx_path = tmp_path / "student.docx"
    _make_docx(sample_resume_text).seek(0)
    Document(_make_docx(sample_resume_text)).save(docx_path)
    pdf_path = tmp_path / "student.pdf"
    pdf_path.write_bytes(_make_pdf(sample_resume_text))

    docx_resume = parse_resume_file(docx_path)
    pdf_resume = parse_resume_file(pdf_path)

    assert docx_resume.profile["email"] == pdf_resume.profile["email"]
    assert "Flask" in docx_resume.profile["skills"]["Backend Development"]
    assert "Flask" in pdf_resume.profile["skills"]["Backend Development"]


def test_scanned_pdf_is_rejected_with_clear_error(client):
    _register(client, "scan@example.com")

    response = client.post(
        "/resumes/",
        data={
            "resume_file": (
                BytesIO(_make_pdf("")),
                "scanned.pdf",
                "application/pdf",
            )
        },
        content_type="multipart/form-data",
        follow_redirects=True,
    )

    assert response.status_code == 200
    assert b"Scanned or image-only PDFs are not supported" in response.data


def test_upload_rejects_unsupported_extension(client):
    _register(client, "extension@example.com")

    response = client.post(
        "/resumes/",
        data={"resume_file": (BytesIO(b"resume text"), "resume.txt", "text/plain")},
        content_type="multipart/form-data",
        follow_redirects=True,
    )

    assert b"Unsupported file type" in response.data


def test_upload_rejects_mismatched_mime_type(client):
    _register(client, "mime@example.com")

    response = client.post(
        "/resumes/",
        data={
            "resume_file": (
                BytesIO(_make_pdf("Resume")),
                "resume.pdf",
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        content_type="multipart/form-data",
        follow_redirects=True,
    )

    assert b"does not match its extension" in response.data


def test_upload_rejects_invalid_magic_bytes(client):
    _register(client, "magic@example.com")

    response = client.post(
        "/resumes/",
        data={
            "resume_file": (BytesIO(b"not a PDF document"), "resume.pdf", "application/pdf")
        },
        content_type="multipart/form-data",
        follow_redirects=True,
    )

    assert b"do not match a valid PDF or DOCX" in response.data


def test_upload_rejects_files_over_five_megabytes(client):
    _register(client, "large@example.com")

    response = client.post(
        "/resumes/",
        data={
            "resume_file": (
                BytesIO(b"%PDF-" + b"x" * (5 * 1024 * 1024)),
                "resume.pdf",
                "application/pdf",
            )
        },
        content_type="multipart/form-data",
        follow_redirects=True,
    )

    assert b"Resume must be 5 MB or smaller" in response.data


def test_multiple_resumes_can_be_saved_and_active_resume_changed(client, app, sample_resume_text):
    _register(client, "multiple@example.com")
    for filename in ("first.docx", "second.docx"):
        client.post(
            "/resumes/",
            data={"resume_file": (_make_docx(sample_resume_text), filename, DOCX_MIME)},
            content_type="multipart/form-data",
        )

    with app.app_context():
        resumes = db.session.scalars(db.select(Resume).order_by(Resume.id)).all()
        assert len(resumes) == 2
        assert [resume.is_active for resume in resumes] == [False, True]
        first_id = resumes[0].id

    response = client.post(f"/resumes/{first_id}/activate", follow_redirects=True)

    assert response.status_code == 200
    assert b"Active resume updated" in response.data
    with app.app_context():
        resumes = db.session.scalars(db.select(Resume).order_by(Resume.id)).all()
        assert [resume.is_active for resume in resumes] == [True, False]


def test_another_user_cannot_view_resume(client, app, sample_resume_text):
    _register(client, "owner@example.com")
    client.post(
        "/resumes/",
        data={"resume_file": (_make_docx(sample_resume_text), "private.docx", DOCX_MIME)},
        content_type="multipart/form-data",
    )
    with app.app_context():
        resume_id = db.session.scalar(db.select(Resume.id))
    client.post("/logout")
    _register(client, "other@example.com")

    assert client.get(f"/resumes/{resume_id}").status_code == 404
    assert client.post(f"/resumes/{resume_id}/activate").status_code == 404


def test_uploaded_original_is_removed_when_parser_rejects_document(client, app):
    _register(client, "empty@example.com")
    document = Document()
    output = BytesIO()
    document.save(output)
    output.seek(0)

    response = client.post(
        "/resumes/",
        data={"resume_file": (output, "empty.docx", DOCX_MIME)},
        content_type="multipart/form-data",
        follow_redirects=True,
    )

    assert b"No readable text was found" in response.data
    assert list(Path(app.config["UPLOAD_FOLDER"]).glob("*")) == []
    with app.app_context():
        assert db.session.scalar(db.select(Resume)) is None
